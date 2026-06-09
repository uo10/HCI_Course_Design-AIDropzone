"""
File-content preview helper — shared by Mock and LLM parsers.

Extracts plain-text previews from text files, PDFs, Office documents, etc.
so the parser can feed actual content (not just filenames) to the AI.

Supports:
    Plain text: txt md csv log json xml yaml py js ts java cpp html css sql ...
    PDF:        pdf  (via PyPDF2)
    Word:       docx (via python-docx)
    Excel:      xlsx (via openpyxl)
    PowerPoint: pptx (via python-pptx)

Public API:
    read_text_preview(path, max_chars=4000) → str
"""

from __future__ import annotations

import codecs
from pathlib import Path

# ---------- extension classification ----------

_TEXT_EXTENSIONS: set[str] = {
    "txt", "md", "csv", "log", "json", "xml", "yaml", "yml",
    "py", "js", "ts", "jsx", "tsx", "java", "cpp", "c", "h",
    "go", "rs", "rb", "php", "swift", "kt", "scala", "lua",
    "html", "css", "scss", "less", "sql", "sh", "bash", "zsh",
    "ps1", "bat", "cmd", "ini", "cfg", "toml", "env", "conf",
    "tex", "rst", "org", "r", "rmd", "ipynb", "dockerfile",
    "gitignore", "makefile", "cmake", "gradle", "lock", "rtf",
}

_DOCUMENT_EXTENSIONS: dict[str, str] = {
    "pdf":  "pdf",
    "docx": "docx",
    "xlsx": "xlsx",
    "pptx": "pptx",
}

# Binary extensions — never attempt to read
_BINARY_EXTENSIONS: set[str] = {
    "jpg", "jpeg", "png", "gif", "bmp", "webp", "ico", "svg",
    "mp4", "avi", "mov", "mkv", "wmv", "flv",
    "mp3", "wav", "flac", "aac", "ogg", "wma",
    "zip", "rar", "7z", "tar", "gz", "bz2", "xz",
    "exe", "dll", "so", "dylib", "obj", "o", "class",
    "pyc", "pyo", "pyd",
}

# ---------- encoding helpers ----------

_FALLBACK_ENCODINGS = ["utf-8", "utf-16", "gbk", "latin-1"]


def _read_as_text(path: Path, max_chars: int) -> str:
    """Read first *max_chars* chars as plain text (multi-encoding fallback)."""
    raw = path.read_bytes()
    for enc in _FALLBACK_ENCODINGS:
        try:
            text = codecs.decode(raw, enc, errors="strict")
            return text[:max_chars] + (
                "\n... (truncated)" if len(text) > max_chars else ""
            )
        except (UnicodeDecodeError, LookupError):
            continue
    return ""


# ---------- document format readers ----------

def _read_pdf(path: Path, max_chars: int) -> str:
    """Extract text from a PDF via PyPDF2 (no system deps)."""
    try:
        from PyPDF2 import PdfReader
        reader = PdfReader(str(path))
        parts: list[str] = []
        for page in reader.pages[:10]:  # first 10 pages
            t = (page.extract_text() or "").strip()
            if t:
                parts.append(t)
            if sum(len(p) for p in parts) >= max_chars:
                break
        text = "\n".join(parts)
        return text[:max_chars] + (
            "\n... (truncated)" if len(text) > max_chars else ""
        )
    except Exception:
        return ""


def _read_docx(path: Path, max_chars: int) -> str:
    """Extract text from a .docx via python-docx."""
    try:
        from docx import Document
        doc = Document(str(path))
        parts: list[str] = []
        for para in doc.paragraphs:
            t = para.text.strip()
            if t:
                parts.append(t)
            if sum(len(p) for p in parts) >= max_chars:
                break
        text = "\n".join(parts)
        return text[:max_chars] + (
            "\n... (truncated)" if len(text) > max_chars else ""
        )
    except Exception:
        return ""


def _read_xlsx(path: Path, max_chars: int) -> str:
    """Extract text from an .xlsx via openpyxl."""
    try:
        from openpyxl import load_workbook
        wb = load_workbook(str(path), read_only=True, data_only=True)
        parts: list[str] = []
        for sheet in wb.worksheets[:3]:  # first 3 sheets
            for row in sheet.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None]
                if cells:
                    parts.append("\t".join(cells))
                if sum(len(p) for p in parts) >= max_chars:
                    break
        wb.close()
        text = "\n".join(parts)
        return text[:max_chars] + (
            "\n... (truncated)" if len(text) > max_chars else ""
        )
    except Exception:
        return ""


def _read_pptx(path: Path, max_chars: int) -> str:
    """Extract text from a .pptx via python-pptx."""
    try:
        from pptx import Presentation
        prs = Presentation(str(path))
        parts: list[str] = []
        for slide in prs.slides[:20]:  # first 20 slides
            for shape in slide.shapes:
                if shape.has_text_frame:
                    t = shape.text_frame.text.strip()
                    if t:
                        parts.append(t)
            if sum(len(p) for p in parts) >= max_chars:
                break
        text = "\n".join(parts)
        return text[:max_chars] + (
            "\n... (truncated)" if len(text) > max_chars else ""
        )
    except Exception:
        return ""


# ---------- dispatch ----------

_DOC_READERS = {
    "pdf":  _read_pdf,
    "docx": _read_docx,
    "xlsx": _read_xlsx,
    "pptx": _read_pptx,
}


def read_text_preview(path: Path, max_chars: int = 4000) -> str:
    """Return the first *max_chars* characters of a file's content.

    Text files are read directly.  PDF/DOCX/XLSX/PPTX are extracted via
    dedicated libraries.  Binary files (images, video, audio) return "".

    If a document library isn't installed, that format returns "" silently.
    """
    try:
        if not path.is_file():
            return ""

        ext = path.suffix.lstrip(".").lower()

        # Document formats
        reader = _DOC_READERS.get(ext)
        if reader:
            return reader(path, max_chars)

        # Plain text formats
        if ext in _TEXT_EXTENSIONS:
            return _read_as_text(path, max_chars)

        # Explicit binary — no attempt
        if ext in _BINARY_EXTENSIONS:
            return ""

        # Unknown extension — probe for text
        try:
            head = path.read_bytes()[:512]
            if b"\x00" in head:
                return ""
            head.decode("utf-8")
            return _read_as_text(path, max_chars)
        except (OSError, UnicodeDecodeError):
            return ""

    except (OSError, PermissionError):
        return ""
